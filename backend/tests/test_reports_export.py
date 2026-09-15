"""DOCX/PDF export for generated reports (Milestone 10 step 4 -- spec
section 4: "DOCX and PDF export: open the generated file again and
assert the headings are there")."""

from io import BytesIO

import docx
import pypdf

from app.modules.reports.export import render_docx, render_pdf

_SECTIONS = [
    {"title": "Abstract", "content": "This project studies poultry disease detection."},
    {"title": "References", "content": "Not covered by your documents."},
]


def test_render_docx_contains_the_title_and_every_section_heading():
    content = render_docx("Project Synopsis", _SECTIONS)

    document = docx.Document(BytesIO(content))
    headings = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style is not None and paragraph.style.name.startswith("Heading")
    ]
    body = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "Project Synopsis" in headings
    assert "Abstract" in headings
    assert "References" in headings
    assert "poultry disease detection" in body


def test_render_pdf_contains_the_title_and_every_section_heading():
    content = render_pdf("Project Synopsis", _SECTIONS)

    reader = pypdf.PdfReader(BytesIO(content))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)

    assert "Project Synopsis" in text
    assert "Abstract" in text
    assert "References" in text
    assert "poultry disease detection" in text


def test_render_docx_and_pdf_agree_on_section_titles():
    """Both formats render from the same `sections` input -- the
    invariant that makes "nothing is stored twice" safe (spec section
    4): render either one from the same data and they cannot disagree."""
    docx_content = render_docx("Study Summary", _SECTIONS)
    pdf_content = render_pdf("Study Summary", _SECTIONS)

    docx_text = "\n".join(p.text for p in docx.Document(BytesIO(docx_content)).paragraphs)
    pdf_text = "\n".join(page.extract_text() or "" for page in pypdf.PdfReader(BytesIO(pdf_content)).pages)

    for section in _SECTIONS:
        assert section["title"] in docx_text
        assert section["title"] in pdf_text
