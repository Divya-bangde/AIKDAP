"""Renders a generated report's stored sections to DOCX or PDF, on
demand, from `Asset.asset_metadata["sections"]`.

Deliberately the only place either format is produced: `service.py`
calls exactly one of these per download request and stores nothing --
so DOCX and PDF can never drift apart, and a report is never rendered
ahead of being asked for (spec section 4).
"""

from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


def render_docx(title: str, sections: list[dict]) -> bytes:
    """Render `title` and every section as a `.docx` file, real headings
    (Word's built-in Heading styles, not bold text) per section."""
    document = Document()
    document.add_heading(title, level=1)
    for section in sections:
        document.add_heading(section["title"], level=2)
        document.add_paragraph(section["content"])

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pdf_safe(text: str) -> str:
    """Escape content for reportlab's mini-HTML `Paragraph` markup, and
    turn newlines into line breaks it understands -- plain text passed
    through unescaped can break rendering on a stray `&`/`<`/`>`."""
    return escape(text).replace("\n", "<br/>")


def render_pdf(title: str, sections: list[dict]) -> bytes:
    """Render `title` and every section as a `.pdf` file, using
    reportlab's Platypus flowables so each section heading is a real
    styled `Paragraph`, not a raw text run."""
    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=LETTER)
    styles = getSampleStyleSheet()

    story = [Paragraph(_pdf_safe(title), styles["Title"]), Spacer(1, 12)]
    for section in sections:
        story.append(Paragraph(_pdf_safe(section["title"]), styles["Heading1"]))
        story.append(Paragraph(_pdf_safe(section["content"]), styles["BodyText"]))
        story.append(Spacer(1, 12))

    document.build(story)
    return buffer.getvalue()
